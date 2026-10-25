import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette - Professional Dark Blue Theme
    BG_DARK = RGBColor(11, 19, 43)        # #0B132B Deep Midnight Navy
    CARD_BG = RGBColor(28, 37, 65)        # #1C2541 Slate Navy Card
    CARD_BORDER = RGBColor(58, 80, 107)   # #3A506B Card outline
    CYAN_ACCENT = RGBColor(56, 189, 248)  # #38BDF8 Sky Cyan Accent
    BLUE_ACCENT = RGBColor(37, 99, 235)   # #2563EB Primary Blue
    GREEN_ACCENT = RGBColor(16, 185, 129) # #10B981 Emerald Green (Verified/Match)
    RED_ACCENT = RGBColor(239, 68, 68)    # #EF4444 Red (Tamper Alert)
    AMBER_ACCENT = RGBColor(245, 158, 11) # #F59E0B Amber (Highlight/PoW)
    TEXT_LIGHT = RGBColor(248, 250, 252)  # #F8FAFC Heading text
    TEXT_MUTED = RGBColor(148, 163, 184)  # #94A3B8 Secondary text
    TEXT_BODY = RGBColor(226, 232, 240)   # #E2E8F0 Readable body text
    INNER_BOX_BG = RGBColor(18, 24, 46)   # #12182E Darker inner container

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5)
        )
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_DARK
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category_text, slide_num_str):
        # Category badge
        badge = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.42), Inches(3.2), Inches(0.36)
        )
        badge.fill.solid()
        badge.fill.fore_color.rgb = BLUE_ACCENT
        badge.line.fill.background()
        tf_b = badge.text_frame
        tf_b.word_wrap = False
        p_b = tf_b.paragraphs[0]
        p_b.text = category_text.upper()
        p_b.font.name = "Segoe UI"
        p_b.font.size = Pt(10)
        p_b.font.bold = True
        p_b.font.color.rgb = TEXT_LIGHT
        p_b.alignment = PP_ALIGN.CENTER

        # Slide Title
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.78), Inches(10.5), Inches(0.8))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.name = "Segoe UI"
        p.font.size = Pt(26)
        p.font.bold = True
        p.font.color.rgb = TEXT_LIGHT

        # Slide Number Badge (Top Right)
        num_box = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(11.8), Inches(0.45), Inches(0.8), Inches(0.36)
        )
        num_box.fill.solid()
        num_box.fill.fore_color.rgb = CARD_BG
        num_box.line.color.rgb = CARD_BORDER
        num_box.line.width = Pt(1)
        tf_n = num_box.text_frame
        p_n = tf_n.paragraphs[0]
        p_n.text = f"{slide_num_str} / 08"
        p_n.font.name = "Segoe UI"
        p_n.font.size = Pt(11)
        p_n.font.bold = True
        p_n.font.color.rgb = CYAN_ACCENT
        p_n.alignment = PP_ALIGN.CENTER

        # Footer divider line
        line = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(7.0), Inches(11.733), Inches(0.02)
        )
        line.fill.solid()
        line.fill.fore_color.rgb = CARD_BORDER
        line.line.fill.background()

        # Footer text
        ft = slide.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(11.733), Inches(0.35))
        tf_f = ft.text_frame
        p_f = tf_f.paragraphs[0]
        p_f.text = "Digital Identity Verification Using Blockchain  •  College Group Project Presentation"
        p_f.font.name = "Segoe UI"
        p_f.font.size = Pt(10)
        p_f.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 1: Title Slide
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1)

    # Decorative top glow bar
    top_glow = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.6), Inches(11.733), Inches(0.06))
    top_glow.fill.solid()
    top_glow.fill.fore_color.rgb = CYAN_ACCENT
    top_glow.line.fill.background()

    # Project Tag Badge
    tag = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.9), Inches(4.5), Inches(0.42))
    tag.fill.solid()
    tag.fill.fore_color.rgb = CARD_BG
    tag.line.color.rgb = CYAN_ACCENT
    tag.line.width = Pt(1)
    tf_tag = tag.text_frame
    p_tag = tf_tag.paragraphs[0]
    p_tag.text = "ACADEMIC CAPSTONE PROJECT • 2024-2025"
    p_tag.font.name = "Segoe UI"
    p_tag.font.size = Pt(11)
    p_tag.font.bold = True
    p_tag.font.color.rgb = CYAN_ACCENT
    p_tag.alignment = PP_ALIGN.CENTER

    # Main Project Title Box
    title_box = slide1.shapes.add_textbox(Inches(0.8), Inches(1.45), Inches(11.733), Inches(1.6))
    tf_t1 = title_box.text_frame
    tf_t1.word_wrap = True
    p_t1 = tf_t1.paragraphs[0]
    p_t1.text = "Digital Identity Verification\nUsing Blockchain"
    p_t1.font.name = "Segoe UI"
    p_t1.font.size = Pt(40)
    p_t1.font.bold = True
    p_t1.font.color.rgb = TEXT_LIGHT
    p_t1.line_spacing = 1.1

    # Subtitle Box
    sub_box = slide1.shapes.add_textbox(Inches(0.8), Inches(3.2), Inches(11.733), Inches(0.7))
    tf_s1 = sub_box.text_frame
    tf_s1.word_wrap = True
    p_s1 = tf_s1.paragraphs[0]
    p_s1.text = "A Tamper-Proof, Zero-PII Credential Verification Platform Powered by SHA-256 Hashing, Proof-of-Work Consensus & FastAPI"
    p_s1.font.name = "Segoe UI"
    p_s1.font.size = Pt(15)
    p_s1.font.color.rgb = CYAN_ACCENT

    # Team Members Card
    team_card = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.1), Inches(7.5), Inches(2.5))
    team_card.fill.solid()
    team_card.fill.fore_color.rgb = CARD_BG
    team_card.line.color.rgb = CARD_BORDER
    team_card.line.width = Pt(1.5)
    tf_team = team_card.text_frame
    tf_team.word_wrap = True
    tf_team.margin_top = Inches(0.2)
    tf_team.margin_left = Inches(0.3)

    p_tm_hdr = tf_team.paragraphs[0]
    p_tm_hdr.text = "GROUP PRESENTATION TEAM"
    p_tm_hdr.font.name = "Segoe UI"
    p_tm_hdr.font.size = Pt(12)
    p_tm_hdr.font.bold = True
    p_tm_hdr.font.color.rgb = CYAN_ACCENT
    p_tm_hdr.space_after = Pt(8)

    members = [
        ("Team Member 1", "Roll No: [Enter Roll No]  •  Frontend & Verification UI"),
        ("Team Member 2", "Roll No: [Enter Roll No]  •  FastAPI Backend & Cryptography"),
        ("Team Member 3", "Roll No: [Enter Roll No]  •  Proof-of-Work Blockchain & Merkle Trees"),
        ("Team Member 4", "Roll No: [Enter Roll No]  •  Database Schema, Testing & Documentation")
    ]
    for name, role in members:
        p_m = tf_team.add_paragraph()
        p_m.text = f"•  {name}  —  {role}"
        p_m.font.name = "Segoe UI"
        p_m.font.size = Pt(11)
        p_m.font.color.rgb = TEXT_BODY
        p_m.space_after = Pt(4)

    # Department & Guide Info Card (Right side)
    guide_card = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.6), Inches(4.1), Inches(3.933), Inches(2.5))
    guide_card.fill.solid()
    guide_card.fill.fore_color.rgb = CARD_BG
    guide_card.line.color.rgb = CARD_BORDER
    guide_card.line.width = Pt(1.5)
    tf_guide = guide_card.text_frame
    tf_guide.word_wrap = True
    tf_guide.margin_top = Inches(0.2)
    tf_guide.margin_left = Inches(0.3)

    p_gd_hdr = tf_guide.paragraphs[0]
    p_gd_hdr.text = "PROJECT GUIDANCE & INSTITUTION"
    p_gd_hdr.font.name = "Segoe UI"
    p_gd_hdr.font.size = Pt(12)
    p_gd_hdr.font.bold = True
    p_gd_hdr.font.color.rgb = CYAN_ACCENT
    p_gd_hdr.space_after = Pt(8)

    guide_lines = [
        "Department of Computer Science & Engineering",
        "Faculty Guide: Prof. [Faculty Guide Name]",
        "Subject: Blockchain Technology & Application",
        "Project Status: Fully Implemented & Tested",
        "Demo Link: http://localhost:5173"
    ]
    for line_text in guide_lines:
        p_g = tf_guide.add_paragraph()
        p_g.text = f"•  {line_text}"
        p_g.font.name = "Segoe UI"
        p_g.font.size = Pt(11)
        p_g.font.color.rgb = TEXT_BODY
        p_g.space_after = Pt(4)

    # =========================================================================
    # SLIDE 2: Introduction and Problem Statement
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2)
    add_header(slide2, "Introduction & Problem Statement", "System Context & Motivation", "02")

    # Left Card: Traditional Identity Systems (Problem)
    left_card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.6), Inches(4.5))
    left_card.fill.solid()
    left_card.fill.fore_color.rgb = CARD_BG
    left_card.line.color.rgb = RED_ACCENT
    left_card.line.width = Pt(1.5)
    tf_l2 = left_card.text_frame
    tf_l2.word_wrap = True
    tf_l2.margin_left = Inches(0.3)
    tf_l2.margin_top = Inches(0.25)
    tf_l2.margin_right = Inches(0.3)

    p_l2_title = tf_l2.paragraphs[0]
    p_l2_title.text = "CURRENT CHALLENGES & VULNERABILITIES"
    p_l2_title.font.name = "Segoe UI"
    p_l2_title.font.size = Pt(13)
    p_l2_title.font.bold = True
    p_l2_title.font.color.rgb = RED_ACCENT
    p_l2_title.space_after = Pt(10)

    prob_points = [
        ("Centralized Vulnerability", "Traditional identity records are kept in central databases that act as single points of failure, vulnerable to cyber attacks and insider tampering."),
        ("Easy Document Forgery", "Digital documents (PDFs, certificates, student IDs) are easily manipulated with graphic tools, making visual inspection unreliable."),
        ("Privacy & PII Leakage", "Storing complete identity credentials publicly risks massive identity theft and violates privacy laws (GDPR/DPDP Act)."),
        ("Slow Verification Cycles", "Organizations take days or weeks contacting issuing authorities manually to confirm authentic credentials.")
    ]
    for title, desc in prob_points:
        p_head = tf_l2.add_paragraph()
        p_head.text = f"•  {title}"
        p_head.font.name = "Segoe UI"
        p_head.font.size = Pt(11.5)
        p_head.font.bold = True
        p_head.font.color.rgb = TEXT_LIGHT
        
        p_desc = tf_l2.add_paragraph()
        p_desc.text = f"    {desc}"
        p_desc.font.name = "Segoe UI"
        p_desc.font.size = Pt(10.5)
        p_desc.font.color.rgb = TEXT_MUTED
        p_desc.space_after = Pt(6)

    # Right Card: The Blockchain Solution
    right_card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.6), Inches(5.733), Inches(4.5))
    right_card.fill.solid()
    right_card.fill.fore_color.rgb = CARD_BG
    right_card.line.color.rgb = GREEN_ACCENT
    right_card.line.width = Pt(1.5)
    tf_r2 = right_card.text_frame
    tf_r2.word_wrap = True
    tf_r2.margin_left = Inches(0.3)
    tf_r2.margin_top = Inches(0.25)
    tf_r2.margin_right = Inches(0.3)

    p_r2_title = tf_r2.paragraphs[0]
    p_r2_title.text = "OUR BLOCKCHAIN-POWERED SOLUTION"
    p_r2_title.font.name = "Segoe UI"
    p_r2_title.font.size = Pt(13)
    p_r2_title.font.bold = True
    p_r2_title.font.color.rgb = GREEN_ACCENT
    p_r2_title.space_after = Pt(10)

    sol_points = [
        ("Zero-PII On-Chain Storage", "No personal documents or names are anchored on the blockchain. Only irreversible 256-bit cryptographic fingerprints (hashes) are recorded."),
        ("Instant Tamper Detection (Avalanche Effect)", "Modifying even a single character in an identity credential alters the entire SHA-256 digest, triggering an immediate Tamper Alert."),
        ("Immutable Proof-of-Work Ledger", "Blocks are permanently linked through previous hashes and PoW mining. Historical records cannot be rewritten or erased."),
        ("Self-Sovereign Real-Time Audit", "Anyone can verify an identity document against the blockchain in milliseconds without depending on intermediaries.")
    ]
    for title, desc in sol_points:
        p_head = tf_r2.add_paragraph()
        p_head.text = f"•  {title}"
        p_head.font.name = "Segoe UI"
        p_head.font.size = Pt(11.5)
        p_head.font.bold = True
        p_head.font.color.rgb = TEXT_LIGHT
        
        p_desc = tf_r2.add_paragraph()
        p_desc.text = f"    {desc}"
        p_desc.font.name = "Segoe UI"
        p_desc.font.size = Pt(10.5)
        p_desc.font.color.rgb = TEXT_MUTED
        p_desc.space_after = Pt(6)

    # Bottom Banner
    banner2 = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.25), Inches(11.733), Inches(0.55))
    banner2.fill.solid()
    banner2.fill.fore_color.rgb = INNER_BOX_BG
    banner2.line.color.rgb = CARD_BORDER
    banner2.line.width = Pt(1)
    tf_b2 = banner2.text_frame
    p_b2 = tf_b2.paragraphs[0]
    p_b2.text = "CORE PARADIGM: Cryptographic verification replaces institutional trust — proving validity without exposing identity data."
    p_b2.font.name = "Segoe UI"
    p_b2.font.size = Pt(11)
    p_b2.font.bold = True
    p_b2.font.color.rgb = CYAN_ACCENT
    p_b2.alignment = PP_ALIGN.CENTER

    # =========================================================================
    # SLIDE 3: Objectives of the Project
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3)
    add_header(slide3, "Project Objectives & Scope", "Project Goals", "03")

    objectives = [
        ("01", "Zero-PII Cryptographic Anchoring",
         "Eliminate data privacy risks by hashing documents using SHA-256 before blockchain anchoring. Raw files remain private and off-chain at all times."),
        ("02", "Custom Proof-of-Work Blockchain Ledger",
         "Develop a transparent, educational blockchain engine featuring difficulty-targeted mining, Merkle roots, previous hash linkage, and chain validation."),
        ("03", "Real-Time Tamper Detection Engine",
         "Deliver instantaneous mathematical audit capabilities comparing computed file hashes against on-chain records with clear Match / Tamper notifications."),
        ("04", "Full-Stack Production-Ready Web Interface",
         "Implement a responsive React 19 single-page app and FastAPI backend with interactive drag-and-drop verification, block explorer, and audit history.")
    ]

    coords = [
        (Inches(0.8), Inches(1.6)),
        (Inches(6.8), Inches(1.6)),
        (Inches(0.8), Inches(4.1)),
        (Inches(6.8), Inches(4.1))
    ]

    for (num, title, desc), (x, y) in zip(objectives, coords):
        card = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.733), Inches(2.25))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER
        card.line.width = Pt(1.5)

        # Number pill inside card
        pill = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x + Inches(0.25), y + Inches(0.25), Inches(0.8), Inches(0.35))
        pill.fill.solid()
        pill.fill.fore_color.rgb = BLUE_ACCENT
        pill.line.fill.background()
        tf_p = pill.text_frame
        p_num = tf_p.paragraphs[0]
        p_num.text = num
        p_num.font.name = "Segoe UI"
        p_num.font.size = Pt(11)
        p_num.font.bold = True
        p_num.font.color.rgb = TEXT_LIGHT
        p_num.alignment = PP_ALIGN.CENTER

        # Title
        tb_t = slide3.shapes.add_textbox(x + Inches(1.2), y + Inches(0.18), Inches(4.3), Inches(0.5))
        tf_t = tb_t.text_frame
        tf_t.word_wrap = True
        p_t = tf_t.paragraphs[0]
        p_t.text = title
        p_t.font.name = "Segoe UI"
        p_t.font.size = Pt(14)
        p_t.font.bold = True
        p_t.font.color.rgb = CYAN_ACCENT

        # Description
        tb_d = slide3.shapes.add_textbox(x + Inches(0.25), y + Inches(0.75), Inches(5.2), Inches(1.3))
        tf_d = tb_d.text_frame
        tf_d.word_wrap = True
        p_d = tf_d.paragraphs[0]
        p_d.text = desc
        p_d.font.name = "Segoe UI"
        p_d.font.size = Pt(11.5)
        p_d.font.color.rgb = TEXT_BODY
        p_d.line_spacing = 1.2

    # Bottom summary tag
    banner3 = slide3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.5), Inches(11.733), Inches(0.35))
    banner3.fill.solid()
    banner3.fill.fore_color.rgb = INNER_BOX_BG
    banner3.line.fill.background()
    tf_b3 = banner3.text_frame
    p_b3 = tf_b3.paragraphs[0]
    p_b3.text = "Target Audience: Academic institutions, employers, certification authorities, and regulatory bodies."
    p_b3.font.name = "Segoe UI"
    p_b3.font.size = Pt(10.5)
    p_b3.font.color.rgb = TEXT_MUTED
    p_b3.alignment = PP_ALIGN.CENTER

    # =========================================================================
    # SLIDE 4: System Architecture and Workflow Diagram
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4)
    add_header(slide4, "System Architecture & Workflow Diagram", "Technical Architecture", "04")

    # 4 Main Architectural Tier Boxes
    tiers = [
        ("1. FRONTEND TIER", "React 19 + Vite", CYAN_ACCENT, [
            "• Drag-and-drop file upload zone",
            "• Client-side SHA-256 pre-calculation",
            "• Live side-by-side hash diff viewer",
            "• Blockchain block explorer & audit log"
        ]),
        ("2. REST API TIER", "FastAPI + Python 3.11", BLUE_ACCENT, [
            "• High-performance async REST routes",
            "• In-memory cryptographic hashing",
            "• Input validation & sanitization",
            "• Auto-generated Swagger documentation"
        ]),
        ("3. OFF-CHAIN STORAGE", "SQLite + SQLAlchemy", AMBER_ACCENT, [
            "• IdentityRecord: ID, name, email",
            "• VerificationLog: Full audit history",
            "• Stores transaction references & blocks",
            "• Zero raw document bytes stored"
        ]),
        ("4. IMMUTABLE LEDGER", "Simulated PoW / Ganache", GREEN_ACCENT, [
            "• SHA-256 Block Header chaining",
            "• Proof of Work mining (difficulty=2)",
            "• Merkle root tree validation",
            "• State trie mapping ID -> Doc Hash"
        ])
    ]

    tier_width = Inches(2.7)
    gap = Inches(0.3)
    start_x = Inches(0.8)

    for i, (t_title, t_sub, color, bullet_list) in enumerate(tiers):
        tx = start_x + i * (tier_width + gap)
        box = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, tx, Inches(1.5), tier_width, Inches(3.0))
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_BG
        box.line.color.rgb = color
        box.line.width = Pt(1.5)

        tf_box = box.text_frame
        tf_box.word_wrap = True
        tf_box.margin_left = Inches(0.18)
        tf_box.margin_right = Inches(0.18)
        tf_box.margin_top = Inches(0.18)

        p_h = tf_box.paragraphs[0]
        p_h.text = t_title
        p_h.font.name = "Segoe UI"
        p_h.font.size = Pt(11)
        p_h.font.bold = True
        p_h.font.color.rgb = color

        p_s = tf_box.add_paragraph()
        p_s.text = t_sub
        p_s.font.name = "Segoe UI"
        p_s.font.size = Pt(10)
        p_s.font.bold = True
        p_s.font.color.rgb = TEXT_LIGHT
        p_s.space_after = Pt(8)

        for bullet in bullet_list:
            p_b = tf_box.add_paragraph()
            p_b.text = bullet
            p_b.font.name = "Segoe UI"
            p_b.font.size = Pt(9.5)
            p_b.font.color.rgb = TEXT_BODY
            p_b.space_after = Pt(3)

        # Connector Arrow (except last box)
        if i < 3:
            arrow = slide4.shapes.add_shape(
                MSO_SHAPE.RIGHT_ARROW, tx + tier_width + Inches(0.06), Inches(2.8), Inches(0.18), Inches(0.25)
            )
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = CYAN_ACCENT
            arrow.line.fill.background()

    # Bottom Workflow Box: Registration vs Verification Flows
    flow_box = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.7), Inches(11.733), Inches(2.1))
    flow_box.fill.solid()
    flow_box.fill.fore_color.rgb = INNER_BOX_BG
    flow_box.line.color.rgb = CARD_BORDER
    flow_box.line.width = Pt(1)

    tf_flow = flow_box.text_frame
    tf_flow.word_wrap = True
    tf_flow.margin_left = Inches(0.3)
    tf_flow.margin_top = Inches(0.15)
    tf_flow.margin_right = Inches(0.3)

    p_fl_title = tf_flow.paragraphs[0]
    p_fl_title.text = "CORE SYSTEM WORKFLOWS (DATA PIPELINE)"
    p_fl_title.font.name = "Segoe UI"
    p_fl_title.font.size = Pt(11)
    p_fl_title.font.bold = True
    p_fl_title.font.color.rgb = CYAN_ACCENT
    p_fl_title.space_after = Pt(6)

    p_f1 = tf_flow.add_paragraph()
    p_f1.text = "A. REGISTRATION WORKFLOW:"
    p_f1.font.name = "Segoe UI"
    p_f1.font.size = Pt(10.5)
    p_f1.font.bold = True
    p_f1.font.color.rgb = GREEN_ACCENT

    p_f1_desc = tf_flow.add_paragraph()
    p_f1_desc.text = "   User uploads Document + ID  →  FastAPI hashes file via SHA-256  →  Simulated PoW mines Block (difficulty=2)  →  Block committed to ledger  →  Metadata & Tx Hash stored in SQLite  →  Receipt generated."
    p_f1_desc.font.name = "Segoe UI"
    p_f1_desc.font.size = Pt(10)
    p_f1_desc.font.color.rgb = TEXT_BODY
    p_f1_desc.space_after = Pt(6)

    p_f2 = tf_flow.add_paragraph()
    p_f2.text = "B. VERIFICATION WORKFLOW:"
    p_f2.font.name = "Segoe UI"
    p_f2.font.size = Pt(10.5)
    p_f2.font.bold = True
    p_f2.font.color.rgb = AMBER_ACCENT

    p_f2_desc = tf_flow.add_paragraph()
    p_f2_desc.text = "   Verifier uploads test Document  →  SHA-256 computed on bytes  →  Query registered hash from Blockchain  →  Bitwise comparison  →  Match (Green Confirmed) or Mismatch (Red Tamper Alert)  →  Result logged to History."
    p_f2_desc.font.name = "Segoe UI"
    p_f2_desc.font.size = Pt(10)
    p_f2_desc.font.color.rgb = TEXT_BODY

    # =========================================================================
    # SLIDE 5: Technologies Used
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5)
    add_header(slide5, "Technologies Used in Implementation", "Tech Stack Breakdown", "05")

    tech_categories = [
        ("FRONTEND DEVELOPMENT", CYAN_ACCENT, [
            ("React 19 & Vite", "High-performance reactive frontend framework with lightning-fast Hot Module Replacement (HMR)."),
            ("Tailwind CSS & Glassmorphism", "Modern responsive styling with custom glassmorphism, glowing status badges, and tab routing."),
            ("Client-Side Hashing", "Web Crypto API implementation for instant client-side SHA-256 checksum previews before upload.")
        ]),
        ("BACKEND & API SERVICES", BLUE_ACCENT, [
            ("FastAPI (Python 3.11)", "Modern, asynchronous web framework delivering high-speed execution and native request typing."),
            ("Uvicorn ASGI Server", "Production-grade asynchronous web server gateway for handling concurrent requests smoothly."),
            ("Pydantic & Multi-Part Form", "Strict schema validation and robust streaming handlers for multipart identity file uploads.")
        ]),
        ("CRYPTOGRAPHY & BLOCKCHAIN", GREEN_ACCENT, [
            ("SHA-256 (Secure Hash Algorithm)", "Cryptographic hash function converting arbitrary document files into fixed 256-bit hexadecimal digests."),
            ("Simulated PoW Blockchain", "Pure Python blockchain ledger featuring difficulty=2 mining, Merkle tree root hashing, and integrity checks."),
            ("Avalanche Effect Protection", "Cryptographic guarantee ensuring even single-bit edits create entirely distinct block digests.")
        ]),
        ("DATABASE & SMART CONTRACTS", AMBER_ACCENT, [
            ("SQLite with SQLAlchemy ORM", "Lightweight relational database managing off-chain metadata (IdentityRecord and VerificationLog)."),
            ("Dual-Mode EVM Architecture", "Built-in support for switching between the simulated ledger and local Ethereum Ganache node."),
            ("Solidity Smart Contract", "Pre-compiled `IdentityVerification.sol` smart contract deployable to local EVM testnet via Web3.py.")
        ])
    ]

    t_coords = [
        (Inches(0.8), Inches(1.55)),
        (Inches(6.8), Inches(1.55)),
        (Inches(0.8), Inches(4.2)),
        (Inches(6.8), Inches(4.2))
    ]

    for (cat_title, color, items), (x, y) in zip(tech_categories, t_coords):
        card = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.733), Inches(2.45))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = color
        card.line.width = Pt(1.5)

        tf_c = card.text_frame
        tf_c.word_wrap = True
        tf_c.margin_left = Inches(0.25)
        tf_c.margin_right = Inches(0.25)
        tf_c.margin_top = Inches(0.2)

        p_h = tf_c.paragraphs[0]
        p_h.text = cat_title
        p_h.font.name = "Segoe UI"
        p_h.font.size = Pt(12)
        p_h.font.bold = True
        p_h.font.color.rgb = color
        p_h.space_after = Pt(6)

        for item_name, item_desc in items:
            p_i = tf_c.add_paragraph()
            p_i.text = f"•  {item_name}: "
            p_i.font.name = "Segoe UI"
            p_i.font.size = Pt(10.5)
            p_i.font.bold = True
            p_i.font.color.rgb = TEXT_LIGHT

            # add desc
            run = p_i.add_run()
            run.text = item_desc
            run.font.bold = False
            run.font.color.rgb = TEXT_MUTED
            p_i.space_after = Pt(4)

    # Bottom note
    banner5 = slide5.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.75), Inches(11.733), Inches(0.2))
    banner5.fill.solid()
    banner5.fill.fore_color.rgb = BG_DARK
    banner5.line.fill.background()

    # =========================================================================
    # SLIDE 6: Working Process
    # =========================================================================
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide6)
    add_header(slide6, "Working Process: Step-by-Step Flow", "System Execution", "06")

    steps = [
        ("STEP 01", "Document Upload & Sanitization", CYAN_ACCENT, [
            "Identity holder submits credential file (.txt, .pdf, .png).",
            "Unique identity ID assigned (e.g., STU-2024-8841).",
            "System validates file stream, rejects empty (0-byte) files.",
            "Raw document is never broadcasted to the blockchain."
        ]),
        ("STEP 02", "SHA-256 Hash Generation", BLUE_ACCENT, [
            "Binary content processed through SHA-256 algorithm.",
            "Produces a unique 64-character hexadecimal digest.",
            "One-way function: impossible to reverse hash back to file.",
            "Deterministic: identical file always yields exact same hash."
        ]),
        ("STEP 03", "Blockchain Record Commitment", AMBER_ACCENT, [
            "Identity ID and Hash combined into a transaction.",
            "Proof-of-Work mining solves target difficulty (e.g. '00...').",
            "Transactions aggregated into block via Merkle root.",
            "Block chained to previous block via cryptographic hash."
        ]),
        ("STEP 04", "Cryptographic Hash Comparison", GREEN_ACCENT, [
            "Verifier uploads document during audit request.",
            "Real-time SHA-256 computed on uploaded document.",
            "Queried against immutable registered hash on the block.",
            "Binary result: 100% Match (Valid) or Tamper Alert (Modified)."
        ])
    ]

    s_width = Inches(2.7)
    s_gap = Inches(0.3)
    s_start_x = Inches(0.8)

    for i, (step_num, step_title, color, step_bullets) in enumerate(steps):
        sx = s_start_x + i * (s_width + s_gap)
        scard = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, sx, Inches(1.5), s_width, Inches(3.8))
        scard.fill.solid()
        scard.fill.fore_color.rgb = CARD_BG
        scard.line.color.rgb = color
        scard.line.width = Pt(1.5)

        tf_s = scard.text_frame
        tf_s.word_wrap = True
        tf_s.margin_left = Inches(0.18)
        tf_s.margin_right = Inches(0.18)
        tf_s.margin_top = Inches(0.18)

        # Step Pill
        p_sp = tf_s.paragraphs[0]
        p_sp.text = step_num
        p_sp.font.name = "Segoe UI"
        p_sp.font.size = Pt(11)
        p_sp.font.bold = True
        p_sp.font.color.rgb = color

        p_st = tf_s.add_paragraph()
        p_st.text = step_title
        p_st.font.name = "Segoe UI"
        p_st.font.size = Pt(12)
        p_st.font.bold = True
        p_st.font.color.rgb = TEXT_LIGHT
        p_st.space_after = Pt(8)

        for b in step_bullets:
            p_sb = tf_s.add_paragraph()
            p_sb.text = f"•  {b}"
            p_sb.font.name = "Segoe UI"
            p_sb.font.size = Pt(10)
            p_sb.font.color.rgb = TEXT_BODY
            p_sb.space_after = Pt(4)

        if i < 3:
            s_arrow = slide6.shapes.add_shape(
                MSO_SHAPE.RIGHT_ARROW, sx + s_width + Inches(0.06), Inches(3.2), Inches(0.18), Inches(0.25)
            )
            s_arrow.fill.solid()
            s_arrow.fill.fore_color.rgb = CYAN_ACCENT
            s_arrow.line.fill.background()

    # Mathematical Equation / Hash Comparison Box Below
    math_box = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.45), Inches(11.733), Inches(1.4))
    math_box.fill.solid()
    math_box.fill.fore_color.rgb = INNER_BOX_BG
    math_box.line.color.rgb = CARD_BORDER
    math_box.line.width = Pt(1)

    tf_m = math_box.text_frame
    tf_m.word_wrap = True
    tf_m.margin_left = Inches(0.3)
    tf_m.margin_top = Inches(0.15)
    tf_m.margin_right = Inches(0.3)

    p_m_title = tf_m.paragraphs[0]
    p_m_title.text = "MATHEMATICAL LOGIC OF VERIFICATION"
    p_m_title.font.name = "Segoe UI"
    p_m_title.font.size = Pt(11)
    p_m_title.font.bold = True
    p_m_title.font.color.rgb = CYAN_ACCENT
    p_m_title.space_after = Pt(4)

    p_eq1 = tf_m.add_paragraph()
    p_eq1.text = "   Given Original Document D and Registered Hash H_reg = SHA-256(D) committed in Block B."
    p_eq1.font.name = "Segoe UI"
    p_eq1.font.size = Pt(10)
    p_eq1.font.color.rgb = TEXT_BODY

    p_eq2 = tf_m.add_paragraph()
    p_eq2.text = "   Upon Audit with Document D' :  If SHA-256(D') == H_reg  →  AUTHENTIC (Confidence = 100%)"
    p_eq2.font.name = "Segoe UI"
    p_eq2.font.size = Pt(10.5)
    p_eq2.font.bold = True
    p_eq2.font.color.rgb = GREEN_ACCENT

    p_eq3 = tf_m.add_paragraph()
    p_eq3.text = "                                 If SHA-256(D') != H_reg  →  TAMPERED (Hash Discrepancy Detected)"
    p_eq3.font.name = "Segoe UI"
    p_eq3.font.size = Pt(10.5)
    p_eq3.font.bold = True
    p_eq3.font.color.rgb = RED_ACCENT

    # =========================================================================
    # SLIDE 7: Live Demonstration & Limitations
    # =========================================================================
    slide7 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide7)
    add_header(slide7, "Live Demonstration, Test Results & Limitations", "Testing & Real-World Scope", "07")

    # Left: Implemented Demo Test Scenarios
    demo_card = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.55), Inches(5.733), Inches(4.5))
    demo_card.fill.solid()
    demo_card.fill.fore_color.rgb = CARD_BG
    demo_card.line.color.rgb = GREEN_ACCENT
    demo_card.line.width = Pt(1.5)

    tf_d7 = demo_card.text_frame
    tf_d7.word_wrap = True
    tf_d7.margin_left = Inches(0.25)
    tf_d7.margin_right = Inches(0.25)
    tf_d7.margin_top = Inches(0.2)

    p_d7_title = tf_d7.paragraphs[0]
    p_d7_title.text = "VERIFIED DEMO SCENARIOS (IMPLEMENTED)"
    p_d7_title.font.name = "Segoe UI"
    p_d7_title.font.size = Pt(12)
    p_d7_title.font.bold = True
    p_d7_title.font.color.rgb = GREEN_ACCENT
    p_d7_title.space_after = Pt(8)

    demo_cases = [
        ("1. Authentic Document Test (Aarav Sharma - STU-2024-8841)",
         "Uploaded `aarav_student_id.txt`.\nResult: Computed SHA-256 matches on-chain record on Block #1.\nUI Display: Glowing Green '100% CRYPTOGRAPHIC MATCH'."),
        ("2. Tampered Document Test (1-Byte Modification)",
         "Uploaded `tampered_aarav_student_id.txt` (only 1 character changed).\nResult: SHA-256 hash completely diverged (Avalanche Effect).\nUI Display: Red 'TAMPER ALERT DETECTED' with hash discrepancy."),
        ("3. Blockchain Explorer & Chain Integrity Audit",
         "Genesis Block #0 + mined blocks displayed with Merkle root & nonce.\nClicking 'Re-Audit Chain Integrity' validates all block links via `is_chain_valid()`."),
        ("4. Automated Backend Test Suite",
         "8 out of 8 test cases pass successfully (100% Green) including input sanitization, duplicate ID rejection, and hash sensitivity.")
    ]
    for case_title, case_desc in demo_cases:
        p_ct = tf_d7.add_paragraph()
        p_ct.text = f"• {case_title}"
        p_ct.font.name = "Segoe UI"
        p_ct.font.size = Pt(10.5)
        p_ct.font.bold = True
        p_ct.font.color.rgb = TEXT_LIGHT

        for line in case_desc.split("\n"):
            p_cd = tf_d7.add_paragraph()
            p_cd.text = f"    {line}"
            p_cd.font.name = "Segoe UI"
            p_cd.font.size = Pt(9.5)
            p_cd.font.color.rgb = TEXT_BODY
        p_cd.space_after = Pt(5)

    # Right: Realistic System Limitations
    limit_card = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.55), Inches(5.733), Inches(4.5))
    limit_card.fill.solid()
    limit_card.fill.fore_color.rgb = CARD_BG
    limit_card.line.color.rgb = AMBER_ACCENT
    limit_card.line.width = Pt(1.5)

    tf_l7 = limit_card.text_frame
    tf_l7.word_wrap = True
    tf_l7.margin_left = Inches(0.25)
    tf_l7.margin_right = Inches(0.25)
    tf_l7.margin_top = Inches(0.2)

    p_l7_title = tf_l7.paragraphs[0]
    p_l7_title.text = "CURRENT SYSTEM LIMITATIONS (HONEST SCOPE)"
    p_l7_title.font.name = "Segoe UI"
    p_l7_title.font.size = Pt(12)
    p_l7_title.font.bold = True
    p_l7_title.font.color.rgb = AMBER_ACCENT
    p_l7_title.space_after = Pt(8)

    limitations = [
        ("Simulated Ledger Persistence",
         "The simulated blockchain is in-memory by default; metadata is saved in SQLite, but the block tree resets on server restart unless re-seeded."),
        ("Private Ganache EVM Environment",
         "Smart contract integration is demonstrated using a local Ganache testnet node. The project is intentionally not connected to public Ethereum mainnet to avoid gas costs and latency."),
        ("'Garbage In, Garbage Out' Boundary",
         "The blockchain cryptographically proves a document has not been altered since registration. However, it cannot verify if the initial document presented at registration was forged without institutional digital signatures."),
        ("Centralized Web Server Deployment",
         "The FastAPI backend and SQLite database currently run on a single host. A full production implementation requires decentralized storage (e.g., IPFS) and distributed nodes.")
    ]
    for lim_title, lim_desc in limitations:
        p_lt = tf_l7.add_paragraph()
        p_lt.text = f"• {lim_title}"
        p_lt.font.name = "Segoe UI"
        p_lt.font.size = Pt(10.5)
        p_lt.font.bold = True
        p_lt.font.color.rgb = TEXT_LIGHT

        p_ld = tf_l7.add_paragraph()
        p_ld.text = f"    {lim_desc}"
        p_ld.font.name = "Segoe UI"
        p_ld.font.size = Pt(9.5)
        p_ld.font.color.rgb = TEXT_MUTED
        p_ld.space_after = Pt(5)

    # Bottom Callout
    banner7 = slide7.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.2), Inches(11.733), Inches(0.55))
    banner7.fill.solid()
    banner7.fill.fore_color.rgb = INNER_BOX_BG
    banner7.line.color.rgb = CARD_BORDER
    banner7.line.width = Pt(1)
    tf_b7 = banner7.text_frame
    p_b7 = tf_b7.paragraphs[0]
    p_b7.text = "KEY TAKEAWAY: The system delivers 100% mathematically verifiable tamper detection within its demonstrated scope."
    p_b7.font.name = "Segoe UI"
    p_b7.font.size = Pt(11)
    p_b7.font.bold = True
    p_b7.font.color.rgb = CYAN_ACCENT
    p_b7.alignment = PP_ALIGN.CENTER

    # =========================================================================
    # SLIDE 8: Conclusion & Future Enhancements
    # =========================================================================
    slide8 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide8)
    add_header(slide8, "Conclusion & Future Roadmap", "Project Wrap-Up", "08")

    # Left: Conclusion
    c_card = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.55), Inches(5.733), Inches(4.5))
    c_card.fill.solid()
    c_card.fill.fore_color.rgb = CARD_BG
    c_card.line.color.rgb = CYAN_ACCENT
    c_card.line.width = Pt(1.5)

    tf_c8 = c_card.text_frame
    tf_c8.word_wrap = True
    tf_c8.margin_left = Inches(0.25)
    tf_c8.margin_right = Inches(0.25)
    tf_c8.margin_top = Inches(0.2)

    p_c8_title = tf_c8.paragraphs[0]
    p_c8_title.text = "PROJECT ACHIEVEMENTS & SUMMARY"
    p_c8_title.font.name = "Segoe UI"
    p_c8_title.font.size = Pt(12)
    p_c8_title.font.bold = True
    p_c8_title.font.color.rgb = CYAN_ACCENT
    p_c8_title.space_after = Pt(8)

    conclusions = [
        ("Privacy By Design", "Demonstrated that digital identity verification does not require compromising sensitive personal data on-chain. Zero PII storage model preserves complete user privacy."),
        ("Uncompromising Tamper Detection", "Leveraged SHA-256 cryptographic sensitivity to detect any document alteration—down to a single space or punctuation mark."),
        ("Working End-to-End System", "Engineered and validated a complete full-stack platform combining React 19, FastAPI, SQLite, and an educational Proof-of-Work blockchain ledger."),
        ("100% Test Validation", "Verified system integrity through 8 automated tests covering hashing determinism, mining validity, input guards, and duplicate handling.")
    ]
    for c_title, c_desc in conclusions:
        p_head = tf_c8.add_paragraph()
        p_head.text = f"•  {c_title}"
        p_head.font.name = "Segoe UI"
        p_head.font.size = Pt(11)
        p_head.font.bold = True
        p_head.font.color.rgb = TEXT_LIGHT

        p_body = tf_c8.add_paragraph()
        p_body.text = f"    {c_desc}"
        p_body.font.name = "Segoe UI"
        p_body.font.size = Pt(10)
        p_body.font.color.rgb = TEXT_BODY
        p_body.space_after = Pt(5)

    # Right: Future Enhancements
    f_card = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.55), Inches(5.733), Inches(4.5))
    f_card.fill.solid()
    f_card.fill.fore_color.rgb = CARD_BG
    f_card.line.color.rgb = GREEN_ACCENT
    f_card.line.width = Pt(1.5)

    tf_f8 = f_card.text_frame
    tf_f8.word_wrap = True
    tf_f8.margin_left = Inches(0.25)
    tf_f8.margin_right = Inches(0.25)
    tf_f8.margin_top = Inches(0.2)

    p_f8_title = tf_f8.paragraphs[0]
    p_f8_title.text = "FUTURE ENHANCEMENTS & ROADMAP"
    p_f8_title.font.name = "Segoe UI"
    p_f8_title.font.size = Pt(12)
    p_f8_title.font.bold = True
    p_f8_title.font.color.rgb = GREEN_ACCENT
    p_f8_title.space_after = Pt(8)

    enhancements = [
        ("Decentralized Storage (IPFS Integration)",
         "Integrate InterPlanetary File System (IPFS) to store client-side encrypted identity credentials, storing only cryptographic Content IDs (CIDs) on-chain."),
        ("Zero-Knowledge Proofs (zk-SNARKs)",
         "Enable credential attribute validation (e.g. 'Age >= 18' or 'Valid Degree') without revealing holder identity, birth date, or raw document content."),
        ("Multi-Signature Institutional Issuers",
         "Implement multi-signature smart contracts requiring digital cryptographic signatures from accredited universities or government authorities."),
        ("Public Layer-2 Deployment (Polygon / Arbitrum)",
         "Deploy EVM contracts to public Layer-2 testnets to provide global decentralization with near-zero transaction fees and instant finality.")
    ]
    for f_title, f_desc in enhancements:
        p_head = tf_f8.add_paragraph()
        p_head.text = f"•  {f_title}"
        p_head.font.name = "Segoe UI"
        p_head.font.size = Pt(11)
        p_head.font.bold = True
        p_head.font.color.rgb = TEXT_LIGHT

        p_body = tf_f8.add_paragraph()
        p_body.text = f"    {f_desc}"
        p_body.font.name = "Segoe UI"
        p_body.font.size = Pt(10)
        p_body.font.color.rgb = TEXT_MUTED
        p_body.space_after = Pt(5)

    # Bottom Thank You Banner
    banner8 = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.2), Inches(11.733), Inches(0.55))
    banner8.fill.solid()
    banner8.fill.fore_color.rgb = BLUE_ACCENT
    banner8.line.fill.background()
    tf_b8 = banner8.text_frame
    p_b8 = tf_b8.paragraphs[0]
    p_b8.text = "THANK YOU! QUESTIONS & DISCUSSION WELCOMED."
    p_b8.font.name = "Segoe UI"
    p_b8.font.size = Pt(13)
    p_b8.font.bold = True
    p_b8.font.color.rgb = TEXT_LIGHT
    p_b8.alignment = PP_ALIGN.CENTER

    # Save presentation
    output_path = os.path.abspath("Digital_Identity_Verification.pptx")
    prs.save(output_path)
    print(f"SUCCESS: Presentation saved at: {output_path}")

if __name__ == "__main__":
    create_deck()
